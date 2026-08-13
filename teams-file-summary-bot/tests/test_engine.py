import shutil
from pathlib import Path

import pytest

from core.engine import run as run_engine

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_files"


def test_engine_end_to_end(tmp_path):
    output_path = tmp_path / "report.docx"
    result = run_engine(
        sources=[SAMPLE_DIR],
        query="summarize sales performance by region",
        output_path=output_path,
    )

    assert output_path.exists()
    assert result.files_used == 2
    assert result.files_considered == 2
    assert not result.warnings

    import docx

    document = docx.Document(str(output_path))
    assert len(document.tables) >= 1
    assert len(document.inline_shapes) >= 1
    all_text = "\n".join(p.text for p in document.paragraphs)
    assert "File Summary Report" in all_text
    assert "Executive Summary" in all_text
    assert "Statistics" in all_text


def test_engine_empty_source_raises(tmp_path):
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()
    with pytest.raises(ValueError):
        run_engine(sources=[empty_dir], query="", output_path=tmp_path / "out.docx")


def test_query_filters_to_matching_files(tmp_path):
    only_text_dir = tmp_path / "src"
    only_text_dir.mkdir()
    (only_text_dir / "cats.txt").write_text("Cats are wonderful pets. They sleep a lot. Cats purr when happy.")
    (only_text_dir / "dogs.txt").write_text("Dogs are loyal companions. They love to play fetch. Dogs bark often.")

    output_path = tmp_path / "report.docx"
    result = run_engine(sources=[only_text_dir], query="dogs fetch", output_path=output_path)

    assert result.files_considered == 2
    assert result.files_used == 1
