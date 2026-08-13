"""Assembles the final .docx report: executive summary, per-file sections,
statistics tables, and embedded charts."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core.stats import AggregateStats


def build_report(
    *,
    output_path: Path,
    query: str,
    aggregate: AggregateStats,
    file_summaries: dict[str, list[str]],
    charts: dict[str, bytes],
) -> Path:
    document = Document()

    _add_title_page(document, query, aggregate)
    _add_executive_summary(document, query, aggregate, file_summaries, charts)
    _add_statistics_section(document, aggregate, charts)
    _add_per_file_sections(document, aggregate, file_summaries)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def _add_title_page(document: Document, query: str, aggregate: AggregateStats) -> None:
    title = document.add_heading("File Summary Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if query.strip():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'Request: "{query.strip()}"')
        run.italic = True

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
        f"{aggregate.file_count} file(s) analyzed | "
        f"{aggregate.total_word_count:,} words processed"
    ).font.size = Pt(10)
    document.add_page_break()


def _add_executive_summary(
    document: Document,
    query: str,
    aggregate: AggregateStats,
    file_summaries: dict[str, list[str]],
    charts: dict[str, bytes],
) -> None:
    document.add_heading("Executive Summary", level=1)

    intro = (
        f"This report summarizes {aggregate.file_count} file(s) totaling "
        f"{aggregate.total_word_count:,} words"
    )
    if aggregate.total_table_count:
        intro += f" and {aggregate.total_table_count} data table(s)"
    if query.strip():
        intro += f", focused on your request: \"{query.strip()}\""
    intro += "."
    document.add_paragraph(intro)

    combined_sentences: list[str] = []
    for sentences in file_summaries.values():
        combined_sentences.extend(sentences[:2])
    for sentence in combined_sentences[:8]:
        document.add_paragraph(sentence, style="List Bullet")

    if "volume" in charts:
        document.add_picture(io.BytesIO(charts["volume"]), width=Inches(5.5))


def _add_statistics_section(document: Document, aggregate: AggregateStats, charts: dict[str, bytes]) -> None:
    document.add_heading("Statistics", level=1)

    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "File", "Words", "Sentences", "Tables"
    for fs in aggregate.per_file:
        row = table.add_row().cells
        row[0].text = fs.file_name
        row[1].text = f"{fs.word_count:,}"
        row[2].text = str(fs.sentence_count)
        row[3].text = str(fs.table_count)

    if "keywords" in charts:
        document.add_paragraph()
        document.add_picture(io.BytesIO(charts["keywords"]), width=Inches(5.5))

    numeric_columns = [nc for fs in aggregate.per_file for nc in fs.numeric_columns]
    if numeric_columns:
        document.add_heading("Numeric Data Found in Tables", level=2)
        num_table = document.add_table(rows=1, cols=6)
        num_table.style = "Light Grid Accent 1"
        hdr = num_table.rows[0].cells
        for i, label in enumerate(["Source", "Column", "Count", "Sum", "Mean", "Min / Max"]):
            hdr[i].text = label
        for nc in numeric_columns:
            row = num_table.add_row().cells
            row[0].text = nc.table_source
            row[1].text = nc.column
            row[2].text = str(nc.count)
            row[3].text = f"{nc.total:,.2f}"
            row[4].text = f"{nc.mean:,.2f}"
            row[5].text = f"{nc.minimum:,.2f} / {nc.maximum:,.2f}"

        if "numeric" in charts:
            document.add_paragraph()
            document.add_picture(io.BytesIO(charts["numeric"]), width=Inches(6.0))


def _add_per_file_sections(document: Document, aggregate: AggregateStats, file_summaries: dict[str, list[str]]) -> None:
    document.add_heading("Per-File Detail", level=1)
    for fs in aggregate.per_file:
        document.add_heading(fs.file_name, level=2)
        sentences = file_summaries.get(fs.file_name, [])
        if sentences:
            for sentence in sentences:
                document.add_paragraph(sentence, style="List Bullet")
        else:
            document.add_paragraph("No extractable text summary for this file.")

        if fs.top_keywords:
            kw_line = ", ".join(f"{w} ({c})" for w, c in fs.top_keywords[:8])
            p = document.add_paragraph()
            p.add_run("Top keywords: ").bold = True
            p.add_run(kw_line)
