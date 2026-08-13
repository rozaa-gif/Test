"""Parsers that turn source files into a common ParsedFile representation.

Supported formats: .txt/.md, .docx, .pdf, .csv, .xlsx/.xls
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf", ".csv", ".xlsx", ".xls"}


@dataclass
class ParsedFile:
    path: Path
    text: str = ""
    tables: list = field(default_factory=list)  # list[pandas.DataFrame]
    error: str | None = None
    is_tabular: bool = False  # True when text is just a table dump (csv/xlsx), not prose

    @property
    def name(self) -> str:
        return self.path.name


def parse_file(path: Path) -> ParsedFile:
    suffix = path.suffix.lower()
    try:
        if suffix in (".txt", ".md"):
            return _parse_text(path)
        if suffix == ".docx":
            return _parse_docx(path)
        if suffix == ".pdf":
            return _parse_pdf(path)
        if suffix == ".csv":
            return _parse_csv(path)
        if suffix in (".xlsx", ".xls"):
            return _parse_excel(path)
        return ParsedFile(path=path, error=f"Unsupported file type: {suffix}")
    except Exception as exc:  # noqa: BLE001 - surface as a parse error, don't crash the run
        return ParsedFile(path=path, error=f"Failed to parse: {exc}")


def _parse_text(path: Path) -> ParsedFile:
    text = path.read_text(encoding="utf-8", errors="replace")
    return ParsedFile(path=path, text=text)


def _parse_docx(path: Path) -> ParsedFile:
    import docx  # python-docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    tables = []
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if rows:
            df = pd.DataFrame(rows[1:], columns=rows[0]) if len(rows) > 1 else pd.DataFrame(rows)
            tables.append(_coerce_numeric(df))
    return ParsedFile(path=path, text="\n".join(paragraphs), tables=tables)


def _parse_pdf(path: Path) -> ParsedFile:
    import pdfplumber

    text_parts: list[str] = []
    tables: list[pd.DataFrame] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
            for raw_table in page.extract_tables() or []:
                if not raw_table or len(raw_table) < 2:
                    continue
                df = pd.DataFrame(raw_table[1:], columns=raw_table[0])
                tables.append(_coerce_numeric(df))
    return ParsedFile(path=path, text="\n".join(text_parts), tables=tables)


def _parse_csv(path: Path) -> ParsedFile:
    with open(path, newline="", encoding="utf-8", errors="replace") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
            sep = dialect.delimiter
        except csv.Error:
            sep = ","
    df = pd.read_csv(path, sep=sep, engine="python")
    df = _coerce_numeric(df)
    text = df.to_string(index=False)
    return ParsedFile(path=path, text=text, tables=[df], is_tabular=True)


def _parse_excel(path: Path) -> ParsedFile:
    sheets = pd.read_excel(path, sheet_name=None)
    tables = [_coerce_numeric(df) for df in sheets.values() if not df.empty]
    text = "\n\n".join(df.to_string(index=False) for df in tables)
    return ParsedFile(path=path, text=text, tables=tables, is_tabular=True)


def _coerce_numeric(df: pd.DataFrame) -> pd.DataFrame:
    """Best-effort conversion of object columns that are actually numeric."""
    for col in df.columns:
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.notna().sum() >= max(1, int(len(df) * 0.6)):
            df[col] = converted
    return df


def discover_files(root: Path, extensions: set[str] | None = None) -> list[Path]:
    """Recursively find supported files under root (root may also be a single file)."""
    exts = extensions or SUPPORTED_EXTENSIONS
    if root.is_file():
        return [root] if root.suffix.lower() in exts else []
    return sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in exts)
